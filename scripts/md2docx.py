"""
把 Markdown 讲稿转换成中文排版的 Word 文档。

用法：
    python scripts/md2docx.py speaker-notes.md paper-speaker-notes.docx

排版：正文微软雅黑 / 西文 Times New Roman，12pt（小四），1.5 倍行距，
首行缩进两字符；支持 #~#### 标题、**加粗**、*斜体*、`行内代码`、
无序/有序列表、引用、管道表格、分隔线。
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
EN_FONT = "Times New Roman"
BODY_SIZE = 12          # 小四
HEADING_SIZES = {1: 20, 2: 16, 3: 14, 4: 12}
HEADING_COLORS = {1: RGBColor(0x1F, 0x3A, 0x5F), 2: RGBColor(0x1F, 0x4E, 0x79),
                  3: RGBColor(0x2E, 0x74, 0xB5), 4: RGBColor(0x40, 0x40, 0x40)}

INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*\s][^*]*\*)")


def set_run_font(run, size=BODY_SIZE, bold=False, italic=False, color=None):
    """西文用 Times New Roman，中文用微软雅黑。"""
    run.font.name = EN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), CN_FONT)


def add_inline(paragraph, text, size=BODY_SIZE, color=None):
    """写入一段文本，并处理 **加粗**、*斜体*、`代码`。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_run_font(paragraph.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            set_run_font(paragraph.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size - 1, color=color)
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_run_font(paragraph.add_run(part[1:-1]), size, italic=True, color=color)
        else:
            set_run_font(paragraph.add_run(part), size, color=color)


def body_paragraph(doc, text, indent=True, space_after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    if indent:
        # 中文首行缩进两字符；firstLineChars 让 Word 按字符数计算
        pf.first_line_indent = Pt(BODY_SIZE * 2)
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")
    add_inline(p, text)
    return p


def heading(doc, level, text):
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(10 if level == 1 else 6)
    pf.keep_with_next = True
    if level == 1:
        p.alignment = 1
    add_inline(p, text, size=HEADING_SIZES[level], color=HEADING_COLORS[level])
    for run in p.runs:
        run.font.bold = True
    return p


def list_item(doc, text, ordered):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.37)
    marker = "• "
    if ordered:
        m = re.match(r"^(\d+)[.)]\s+(.*)$", text)
        marker, text = m.group(1) + ". ", m.group(2)
    set_run_font(p.add_run(marker))
    add_inline(p, text)
    return p


def blockquote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.left_indent = Cm(0.74)
    add_inline(p, text, color=RGBColor(0x59, 0x59, 0x59))
    return p


def table(doc, rows):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            text = row[j] if j < len(row) else ""
            add_inline(p, text, size=11)
            if i == 0:
                for run in p.runs:
                    run.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def split_row(line):
    row = line.strip().strip("|")
    return [c.strip() for c in row.split("|")]


def is_separator_row(line):
    return bool(re.fullmatch(r"\|?[\s:|-]+\|?", line.strip())) and "-" in line


def convert(md_path, docx_path):
    if not os.path.exists(md_path):
        print(f"[error] 找不到文件：{md_path}")
        sys.exit(1)

    doc = Document()

    # A4 页面，中文文档常用页边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.18)

    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CN_FONT)

    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line or line in ("---", "***", "___"):
            i += 1
            continue

        # 管道表格（含分隔行）
        if (line.startswith("|") and i + 1 < len(lines)
                and is_separator_row(lines[i + 1])):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            table(doc, rows)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            heading(doc, len(m.group(1)), m.group(2).strip())
            i += 1
            continue

        if re.match(r"^[-*+]\s+", line):
            list_item(doc, re.sub(r"^[-*+]\s+", "", line), ordered=False)
            i += 1
            continue

        if re.match(r"^\d+[.)]\s+", line):
            list_item(doc, line, ordered=True)
            i += 1
            continue

        if line.startswith(">"):
            blockquote(doc, line.lstrip("> ").strip())
            i += 1
            continue

        body_paragraph(doc, line)
        i += 1

    doc.save(docx_path)
    print(f"[success] 已生成 Word 文档：{docx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("用法: python scripts/md2docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
